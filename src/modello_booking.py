"""
modello_booking.py
==================
Il "cervello" del calcolo, copia fedele in Python del modello che avevi gia'
fatto in HTML/JavaScript (Scenario punteggio Booking v2.0).

Idea centrale (EWMA = media pesata sul tempo):
  - ogni recensione pesa di piu' se e' recente;
  - il peso si DIMEZZA ogni 6 mesi (H = 6):   peso = 0.5 ** (eta_in_mesi / 6)
  - il "punteggio calcolato" e' la media dei voti pesata cosi'.

Poi:
  - pace = media semplice degli ultimi 90 giorni (la qualita' che tieni adesso)
  - rate = quante recensioni prendi al mese (dagli ultimi 180 giorni)
  - proiezione = simuli 24 mesi aggiungendo 'rate' recensioni al voto 'pace'
                 e ricalcoli l'EWMA mese per mese.
"""

from __future__ import annotations

import csv
import io
import math
import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path

GIORNI_MESE = 30.44          # giorni medi in un mese (come MS_M nel tuo codice)
H = 6                        # emivita del peso, in mesi (0.5 ogni 6 mesi)

# Mesi abbreviati in italiano, per le etichette del grafico (es. "set 26").
_MESI_IT = ["gen", "feb", "mar", "apr", "mag", "giu",
            "lug", "ago", "set", "ott", "nov", "dic"]

# Parole chiave per riconoscere le colonne dell'export Booking.
_NEEDLE_DATA = ["data della recensione", "review date", "date"]
_NEEDLE_VOTO = ["punteggio della recensione", "review score", "punteggio", "score"]
_AVOID_VOTO = ["staff", "pulizia", "posizione", "servizi", "comfort", "qualit", "prezzo"]

# Sottoreparti: (parola chiave nella colonna, etichetta leggibile)
_CATEGORIE = [
    ("staff", "Personale · Reception"),
    ("pulizia", "Pulizia · Camere"),
    ("comfort", "Comfort · Manutenzione"),
    ("servizi", "Servizi"),
    ("posizion", "Posizione"),
    ("qualit", "Qualità / prezzo"),
]


@dataclass
class Recensione:
    d: datetime   # data della recensione
    s: float      # voto (1-10)


@dataclass
class Categoria:
    key: str
    label: str
    lista: list[Recensione]


# --------------------------------------------------------------------------
#  Lettura del file CSV scaricato da Booking
# --------------------------------------------------------------------------
def _trova_colonna(cols: list[str], needles: list[str], avoid: list[str] | None = None) -> str | None:
    avoid = avoid or []
    for c in cols:
        l = c.lower()
        if any(n in l for n in needles) and not any(a in l for a in avoid):
            return c
    return None


def _parse_data(s) -> datetime | None:
    """Legge una data in vari formati comuni (ISO, gg/mm/aaaa, ecc.)."""
    if s is None:
        return None
    testo = str(s).strip()
    if not testo:
        return None
    # Provo prima il formato ISO (come fa JavaScript con new Date()).
    try:
        return datetime.fromisoformat(testo.replace(" ", "T")[:19])
    except ValueError:
        pass
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d.%m.%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(testo[:10], fmt)
        except ValueError:
            continue
    return None


def _parse_voto(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    testo = str(v).strip().replace(",", ".")
    m = re.search(r"-?\d+(?:\.\d+)?", testo)
    return float(m.group(0)) if m else None


def _righe_csv(path: Path) -> tuple[list[dict], list[str]]:
    testo = path.read_text(encoding="utf-8-sig", errors="replace")
    prima = testo.splitlines()[0] if testo.splitlines() else ""
    sep = ";" if prima.count(";") > prima.count(",") else ","
    lettore = csv.DictReader(io.StringIO(testo), delimiter=sep)
    righe = [dict(r) for r in lettore]
    cols = lettore.fieldnames or (list(righe[0].keys()) if righe else [])
    return righe, list(cols)


def _righe_xlsx(path: Path) -> tuple[list[dict], list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    cols = [str(c).strip() if c is not None else "" for c in next(it, [])]
    righe = []
    for valori in it:
        righe.append({cols[i]: valori[i] for i in range(min(len(cols), len(valori)))})
    return righe, cols


def carica_recensioni(path: Path) -> tuple[list[Recensione], list[Categoria]]:
    """Legge il file scaricato e restituisce (recensioni, categorie)."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        righe, cols = _righe_xlsx(path)
    else:
        righe, cols = _righe_csv(path)

    if not righe:
        return [], []

    d_col = _trova_colonna(cols, _NEEDLE_DATA) or cols[0]
    s_col = _trova_colonna(cols, _NEEDLE_VOTO, _AVOID_VOTO)

    recensioni: list[Recensione] = []
    for r in righe:
        d = _parse_data(r.get(d_col))
        s = _parse_voto(r.get(s_col)) if s_col else None
        if d is not None and s is not None:
            recensioni.append(Recensione(d, s))
    recensioni.sort(key=lambda x: x.d)

    # Sottoreparti (una colonna per categoria, se presente e con abbastanza dati).
    categorie: list[Categoria] = []
    for key, label in _CATEGORIE:
        col = next((c for c in cols if key in c.lower()), None)
        if not col:
            continue
        lista = []
        for r in righe:
            d = _parse_data(r.get(d_col))
            v = _parse_voto(r.get(col))
            if d is not None and v is not None and v > 0:
                lista.append(Recensione(d, v))
        lista.sort(key=lambda x: x.d)
        if len(lista) >= 10:
            categorie.append(Categoria(key, label, lista))

    return recensioni, categorie


# --------------------------------------------------------------------------
#  Il calcolo vero e proprio (EWMA + finestre + proiezione)
# --------------------------------------------------------------------------
def ewma(lista: list[Recensione], as_of: datetime, h: float = H) -> float:
    """Media pesata: il peso si dimezza ogni 'h' mesi di eta'."""
    num = den = 0.0
    for r in lista:
        eta_mesi = (as_of - r.d).total_seconds() / 86400.0 / GIORNI_MESE
        w = 0.5 ** (eta_mesi / h)
        num += w * r.s
        den += w
    return num / den if den else 0.0


def media_finestra(lista: list[Recensione], giorni: int) -> tuple[float, int] | None:
    """Media semplice delle recensioni negli ultimi 'giorni' (dalla piu' recente)."""
    if not lista:
        return None
    fine = lista[-1].d
    taglio = fine - timedelta(days=giorni)
    dentro = [r.s for r in lista if r.d >= taglio]
    if not dentro:
        return None
    return sum(dentro) / len(dentro), len(dentro)


def _fmt_mese(d: datetime) -> str:
    return f"{_MESI_IT[d.month - 1]} {str(d.year)[2:]}"


def _ultimo_giorno_mese(anno: int, mese_zero: int) -> datetime:
    """Come new Date(anno, mese, 0) in JS: ultimo giorno del mese (mese_zero 0-based)."""
    # mese_zero puo' eccedere 11: normalizzo.
    anno += mese_zero // 12
    mese_zero = mese_zero % 12
    if mese_zero == 0:
        return datetime(anno - 1, 12, 31)
    # primo giorno del mese, meno un giorno
    return datetime(anno, mese_zero + 1, 1) - timedelta(days=1) if mese_zero < 12 else datetime(anno, 12, 31)


@dataclass
class Analisi:
    calc: float                      # punteggio calcolato oggi (EWMA)
    n_totali: int
    pace: float                      # media ultimi 3 mesi
    n3: int
    rate: int                        # recensioni al mese
    w7: tuple[float, int] | None
    w30: tuple[float, int] | None
    trend: str                       # "up" | "down" | "flat"
    storico_labels: list[str] = field(default_factory=list)
    storico_valori: list[float] = field(default_factory=list)
    proiezione_labels: list[str] = field(default_factory=list)
    proiezione_valori: list[float] = field(default_factory=list)
    gradini: list[dict] = field(default_factory=list)      # {target, quando, mesi}
    sottoreparti: list[dict] = field(default_factory=list)


def analizza(recensioni: list[Recensione], categorie: list[Categoria] | None = None) -> Analisi | None:
    """Esegue tutto il modello e restituisce i numeri pronti per il cruscotto."""
    if len(recensioni) < 5:
        return None
    categorie = categorie or []

    fine = recensioni[-1].d
    inizio = recensioni[0].d
    calc = ewma(recensioni, fine)

    w3 = media_finestra(recensioni, 90)
    if not w3 or w3[1] < 5:
        w3 = media_finestra(recensioni, 180)
    pace, n3 = w3
    w7 = media_finestra(recensioni, 7)
    w30 = media_finestra(recensioni, 30)

    recenti_180 = [r for r in recensioni if r.d >= fine - timedelta(days=180)]
    rate = max(1, round(len(recenti_180) / 6))

    trend = "up" if pace > calc + 0.02 else ("down" if pace < calc - 0.02 else "flat")

    # ---- serie storica: EWMA mese per mese ----
    storico_labels, storico_valori = [], []
    cur = _ultimo_giorno_mese(inizio.year, inizio.month)  # month 0-based = inizio.month => ultimo giorno mese di inizio
    while cur <= fine:
        sub = [r for r in recensioni if r.d <= cur]
        if sub:
            storico_labels.append(_fmt_mese(cur))
            storico_valori.append(round(ewma(sub, cur), 3))
        cur = _ultimo_giorno_mese(cur.year, cur.month + 1)

    # ---- proiezione 24 mesi ----
    horizon = 24
    proiezione_labels, proiezione_valori = [], []
    base = list(recensioni)
    pm = _ultimo_giorno_mese(fine.year, fine.month + 1)
    # target: dai gradini sopra 'calc' fino a 'pace'
    gradini = []
    t = math.ceil((calc + 0.001) * 10) / 10
    top = round(pace * 10) / 10
    targets = []
    while t <= top + 1e-9:
        targets.append(round(t, 1))
        t = round(t + 0.1, 1)
    hits: dict[float, dict] = {}

    for i in range(horizon):
        proiezione_labels.append(_fmt_mese(pm))
        meta = datetime(pm.year, pm.month, 15)
        for _ in range(rate):
            base.append(Recensione(meta, pace))
        v = ewma(base, pm)
        proiezione_valori.append(round(v, 3))
        for tg in targets:
            if tg not in hits and v >= tg - 0.05:
                hits[tg] = {"quando": _fmt_mese(pm), "mesi": i + 1}
        pm = _ultimo_giorno_mese(pm.year, pm.month + 1)

    for tg in targets[:6]:
        h = hits.get(tg)
        gradini.append({
            "target": tg,
            "quando": h["quando"] if h else "oltre 2 anni",
            "mesi": h["mesi"] if h else None,
            "badge": _badge(tg),
        })

    # ---- sottoreparti ----
    sottoreparti = []
    for c in categorie:
        sc = _scenario_categoria(c.lista)
        if sc:
            sc["label"] = c.label
            sottoreparti.append(sc)

    return Analisi(
        calc=round(calc, 2), n_totali=len(recensioni),
        pace=round(pace, 2), n3=n3, rate=rate,
        w7=(round(w7[0], 2), w7[1]) if w7 else None,
        w30=(round(w30[0], 2), w30[1]) if w30 else None,
        trend=trend,
        storico_labels=storico_labels, storico_valori=storico_valori,
        proiezione_labels=proiezione_labels, proiezione_valori=proiezione_valori,
        gradini=gradini, sottoreparti=sottoreparti,
    )


def _badge(s: float) -> str:
    if s >= 9.5: return "Eccezionale"
    if s >= 9.0: return "Eccellente"
    if s >= 8.6: return "Favoloso"
    if s >= 8.0: return "Ottimo"
    if s >= 7.0: return "Buono"
    return ""


def _scenario_categoria(lista: list[Recensione]) -> dict | None:
    if len(lista) < 5:
        return None
    fine = lista[-1].d
    calc = ewma(lista, fine)
    w3 = media_finestra(lista, 90)
    if not w3 or w3[1] < 5:
        w3 = media_finestra(lista, 180)
    pace = w3[0]
    rate = max(1, round(len([r for r in lista if r.d >= fine - timedelta(days=180)]) / 6))
    trend = "up" if pace > calc + 0.02 else ("down" if pace < calc - 0.02 else "flat")
    prossimo, quando = None, None
    if trend == "up":
        primo = math.ceil((calc + 0.001) * 10) / 10
        top = round(pace * 10) / 10
        if primo <= top + 1e-9:
            base = list(lista)
            pm = _ultimo_giorno_mese(fine.year, fine.month + 1)
            for i in range(24):
                meta = datetime(pm.year, pm.month, 15)
                for _ in range(rate):
                    base.append(Recensione(meta, pace))
                if ewma(base, pm) >= primo - 0.05:
                    prossimo = primo
                    quando = f"{_fmt_mese(pm)} (~{i+1}m)"
                    break
                pm = _ultimo_giorno_mese(pm.year, pm.month + 1)
    return {
        "calc": round(calc, 2), "pace": round(pace, 2), "trend": trend,
        "prossimo": prossimo, "quando": quando,
    }


def analizza_file(path: Path) -> Analisi | None:
    """Scorciatoia: legge il file e lo analizza in un colpo solo."""
    recensioni, categorie = carica_recensioni(Path(path))
    return analizza(recensioni, categorie)
